import csv
import io
import json
from datetime import datetime
from pathlib import Path
from uuid import uuid4
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes.smvu import SmvuEventIn
from app.core.config import Settings
from app.db.models import Equipment, HistoricalEvent, ImportOutbox, IntegrationRun, utcnow
from app.domain import audit, equipment, smvu
from app.schemas.equipment import EquipmentIn

ALIASES = {
    "ид_канала_данных": "asset_id",
    "ид_объект": "object_id",
    "ид_объекта": "object_id",
    "название_датчика": "name",
    "тип_датчика": "equipment_type",
    "тег_инженерной_системы": "tag",
    "тип_инж_системы": "system_type",
    "ид_события": "event_id",
    "дата": "date",
    "время": "time",
    "значение_датчика": "value",
    "тревожное": "alarm",
}
EVENT_HEADERS = ["ид_события", "ид_канала_данных", "дата", "время", "тревожное", "значение_датчика"]


def staging(settings: Settings) -> Path:
    if settings.dataset_dir is None:
        raise equipment.RegistryError("Хранилище загрузок не настроено.")
    path = settings.dataset_dir / ".imports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _rows(file, filename: str, settings: Settings):
    file.seek(0, 2)
    if file.tell() > settings.upload_max_bytes:
        raise equipment.RegistryError("Файл превышает допустимый размер 256 МБ.")
    file.seek(0)
    if filename.lower().endswith(".csv"):
        stream = io.TextIOWrapper(file, encoding="utf-8-sig", newline="")
        try:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            yield from csv.reader(stream, dialect)
        finally:
            stream.detach()
    elif filename.lower().endswith(".xlsx"):
        try:
            with ZipFile(file) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(item.file_size for item in archive.infolist()) > 268435456
                ):
                    raise equipment.RegistryError("Распакованный XLSX превышает 256 МБ.")
                if any("vbaProject" in item.filename for item in archive.infolist()):
                    raise equipment.RegistryError("Макросы не поддерживаются.")
            file.seek(0)
            book = load_workbook(file, read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    raise equipment.RegistryError("В XLSX должен быть ровно один лист с данными.")
                sheet = book.worksheets[0]
                sheet.reset_dimensions()
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise equipment.RegistryError(
                            "Формулы запрещены: сохраните значения ячеек."
                        )
                    if len(row) > 40:
                        raise equipment.RegistryError("В таблице больше 40 колонок.")
                    yield [cell.value for cell in row]
            finally:
                book.close()
        except (BadZipFile, KeyError, DefusedXmlException, ParseError) as error:
            raise equipment.RegistryError("Файл не является корректным XLSX.") from error
    else:
        raise equipment.RegistryError("Выберите файл CSV (UTF-8) или XLSX.")


def _string(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize(row: dict, kind: str, settings: Settings) -> dict:
    values = {key: value for key, value in row.items() if value is not None and value != ""}
    if kind == "channels":
        values.setdefault("external_id", values.get("asset_id"))
        for key in values:
            if key != "installed_on":
                values[key] = _string(values[key])
        return EquipmentIn.model_validate(values).model_dump(mode="json")
    values["channel_id"] = values.pop("asset_id", values.get("channel_id", ""))
    if "ts" not in values:
        day = values.pop("date", "")
        day = day.date().isoformat() if isinstance(day, datetime) else str(day)
        values["ts"] = day + "T" + str(values.pop("time", "00:00:00"))
    for key in ["event_id", "channel_id", "value"]:
        if key in values:
            values[key] = _string(values[key])
    if isinstance(values.get("alarm"), str):
        values["alarm"] = {"да": True, "нет": False}.get(values["alarm"].lower(), values["alarm"])
    event = SmvuEventIn.model_validate(values)
    if event.ts.tzinfo is None:
        event.ts = event.ts.replace(tzinfo=ZoneInfo(settings.timezone))
    return event.model_dump(mode="json")


def preview(
    session: Session, settings: Settings, file, filename: str, kind: str, actor: str
) -> IntegrationRun:
    run = IntegrationRun(
        id=uuid4().hex,
        kind=kind,
        source=settings.equipment_sync_source,
        status="validated",
        detail="",
    )
    path = staging(settings) / f"{run.id}.jsonl"
    errors: list[dict] = []
    total, rejected = 0, 0
    seen: set[str] = set()
    known = set(session.scalars(select(Equipment.asset_id))) if kind == "journal" else set()
    try:
        rows = iter(_rows(file, filename, settings))
        header = next(rows, None)
        if not header or len(header) > 40:
            raise equipment.RegistryError("Не найдена строка заголовков (не более 40 колонок).")
        header = [ALIASES.get(_string(value).lower(), _string(value).lower()) for value in header]
        required = (
            {"asset_id", "name", "object_id", "equipment_type", "tag", "system_type"}
            if kind == "channels"
            else {"event_id", "value", "alarm"}
        )
        if kind == "journal":
            if not ({"date", "time"} <= set(header) or "ts" in header):
                raise equipment.RegistryError("Нужны колонки дата, время либо ts.")
            if not ({"asset_id", "channel_id"} & set(header)):
                raise equipment.RegistryError("Нет колонки ид_канала_данных.")
        missing = required - set(header)
        allowed = (
            set(EquipmentIn.model_fields)
            if kind == "channels"
            else {"event_id", "asset_id", "channel_id", "date", "time", "ts", "value", "alarm"}
        )
        if missing or len(set(header)) != len(header) or set(header) - allowed:
            raise equipment.RegistryError(
                "Проверьте заголовки по шаблону. Отсутствуют: " + ", ".join(sorted(missing))
            )
        with path.open("w", encoding="utf-8") as output:
            for number, row in enumerate(rows, 2):
                if not any(value is not None and value != "" for value in row):
                    continue
                total += 1
                if total > settings.upload_max_rows:
                    raise equipment.RegistryError(
                        "Превышен лимит 1 000 000 строк. Разделите журнал на части."
                    )
                try:
                    if len(row) > len(header) and any(
                        value not in (None, "") for value in row[len(header) :]
                    ):
                        raise ValueError("Число значений не совпадает с заголовком")
                    values = normalize(dict(zip(header, row, strict=False)), kind, settings)
                    identity = values["asset_id" if kind == "channels" else "event_id"]
                    if identity in seen:
                        raise ValueError("Повторяется идентификатор в файле")
                    seen.add(identity)
                    if kind == "journal" and values["channel_id"] not in known:
                        raise ValueError(
                            "Канал отсутствует в реестре: сначала загрузите справочник"
                        )
                    output.write(json.dumps(values, ensure_ascii=False) + "\n")
                except (ValueError, TypeError) as error:
                    rejected += 1
                    message = (
                        "Некорректные поля: " + ", ".join(str(e["loc"][0]) for e in error.errors())
                        if isinstance(error, ValidationError)
                        else str(error)
                    )
                    if len(errors) < 100:
                        errors.append({"row": number, "message": message})
        if total == 0:
            raise equipment.RegistryError("Файл не содержит строк данных.")
        if rejected:
            run.status = "failed"
    except (ValueError, OSError, csv.Error, UnicodeError) as error:
        run.status = "failed"
        errors.append(
            {
                "row": 0,
                "message": str(error)
                if isinstance(error, equipment.RegistryError)
                else "Не удалось прочитать таблицу. Проверьте формат и кодировку UTF-8.",
            }
        )
    if run.status == "failed":
        path.unlink(missing_ok=True)
    run.result = json.dumps(
        {
            "rows": total,
            "valid": total - rejected,
            "rejected": rejected,
            "accepted": 0,
            "errors": errors,
        },
        ensure_ascii=False,
    )
    run.finished_at = utcnow()
    session.add(run)
    audit.record(
        session,
        actor=actor,
        role="admin",
        action="import.validate",
        resource_type="integration_run",
        resource_id=run.id,
        detail=f"{kind}: {total} строк",
    )
    session.flush()
    return run


def apply_import(
    session: Session, settings: Settings, run: IntegrationRun, actor: str
) -> IntegrationRun:
    equipment.lock_registry(session)
    session.refresh(run, with_for_update=True)
    if run.status == "success":
        return run
    if run.status != "validated":
        raise equipment.RegistryError("Сначала исправьте ошибки и проверьте файл.")
    path = staging(settings) / f"{run.id}.jsonl"
    counts = {"created": 0, "updated": 0, "unchanged": 0}
    journal_path = staging(settings) / f"{run.id}.csv"
    with (
        path.open(encoding="utf-8") as handle,
        journal_path.open("w", encoding="utf-8", newline="") as archive,
    ):
        writer = csv.writer(archive)
        writer.writerow(EVENT_HEADERS)
        while True:
            from itertools import islice

            batch = [json.loads(line) for line in islice(handle, 1000)]
            if not batch:
                break
            if run.kind == "channels":
                change = equipment.upsert(
                    session, [EquipmentIn.model_validate(row) for row in batch], run.source
                )
                for key in counts:
                    counts[key] += change[key]
            else:
                existing = {
                    row.event_id: row
                    for row in session.scalars(
                        select(HistoricalEvent).where(
                            HistoricalEvent.event_id.in_([row["event_id"] for row in batch])
                        )
                    )
                }
                for row in batch:
                    event = SmvuEventIn.model_validate(row)
                    current = existing.get(event.event_id)
                    if current:
                        if any(
                            getattr(current, key) != value
                            for key, value in event.model_dump().items()
                        ):
                            raise equipment.RegistryError(
                                f"Событие {event.event_id} уже записано с другими значениями."
                            )
                        counts["unchanged"] += 1
                        continue
                    session.add(HistoricalEvent(**event.model_dump()))
                    ts = event.ts.astimezone(ZoneInfo(settings.timezone))
                    writer.writerow(
                        [
                            event.event_id,
                            event.channel_id,
                            ts.strftime("%Y-%m-%d"),
                            ts.strftime("%H:%M:%S"),
                            "t" if event.alarm else "f",
                            event.value,
                        ]
                    )
                    counts["created"] += 1
                session.flush()
    if run.kind == "channels":
        equipment.queue_registry(session)
        journal_path.unlink(missing_ok=True)
    elif counts["created"]:
        session.add(
            ImportOutbox(id=run.id, payload=json.dumps({"kind": "journal", "run_id": run.id}))
        )
    run.status = "success"
    run.finished_at = utcnow()
    result = json.loads(run.result)
    result.update(counts, accepted=counts["created"] + counts["updated"])
    run.result = json.dumps(result, ensure_ascii=False)
    audit.record(
        session,
        actor=actor,
        role="admin",
        action="import.apply",
        resource_type="integration_run",
        resource_id=run.id,
        detail=run.result,
    )
    session.flush()
    return run


def publish_pending(session: Session, settings: Settings) -> None:
    if settings.dataset_dir is None and settings.inbox_dir is None:
        return
    equipment.lock_registry(session)
    rows = list(session.scalars(select(ImportOutbox).where(ImportOutbox.delivered_at.is_(None))))
    if not rows:
        return
    dataset_changed = False
    for row in rows:
        payload = json.loads(row.payload)
        if payload["kind"] == "stream":
            if settings.inbox_dir is None:
                continue
            name = f"stream-{row.id}.jsonl"
            if (
                not (settings.inbox_dir / name).exists()
                and not (settings.inbox_dir / "processed" / name).exists()
            ):
                smvu.spool_events(
                    settings.inbox_dir,
                    payload["batch_id"],
                    payload["events"],
                    settings.timezone,
                    delivery_id=row.id,
                )
            row.delivered_at = utcnow()
            continue
        if settings.dataset_dir is None:
            continue
        dataset_changed = True
        if payload["kind"] == "channels":
            equipment.publish_registry(session, settings)
        else:
            source = staging(settings) / f"{payload['run_id']}.csv"
            target = settings.dataset_dir / "uploads" / f"ext-journal-{payload['run_id']}.csv"
            target.parent.mkdir(parents=True, exist_ok=True)
            if source.exists():
                source.replace(target)
            elif not target.exists():
                raise OSError("Файл журнала не найден")
        row.delivered_at = utcnow()
    if dataset_changed and settings.dataset_dir is not None:
        marker = settings.dataset_dir / ".revision"
        temporary = marker.with_suffix(".tmp")
        temporary.write_text(uuid4().hex, encoding="utf-8")
        temporary.replace(marker)
    session.commit()
